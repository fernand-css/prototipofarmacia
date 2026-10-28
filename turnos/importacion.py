"""Importación del padrón de pacientes y sus recetas vigentes desde un archivo de Rayén (spec 04).

`leer_archivo` convierte un `.csv` o `.xlsx` en filas de texto. Un archivo que no se puede usar
entero (extensión, tamaño, columnas) lanza `ErrorDeArchivo`, y su mensaje se muestra tal cual a
la QF. `validar_filas` separa las filas válidas, ya normalizadas, de las rechazadas con su motivo.
`preparar_carga` compara esas filas con la base, sin escribir pacientes ni recetas, y deja la
carga pendiente de confirmar. `confirmar_carga` la aplica en una transacción; una acción que no
corresponde sobre una carga lanza `ErrorDeCarga`.
"""
import csv
import hashlib
import io
import re
import unicodedata
from collections import defaultdict
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from pathlib import Path

from django.conf import settings
from django.core.exceptions import ValidationError
from django.db import transaction
from django.db.models import F, Value
from django.db.models.functions import Greatest
from django.utils import timezone

from cuentas.validators import normalizar_rut, validar_rut
from turnos.models import CargaPadron, EstadoCarga, Paciente, RecetaCronica

COLUMNAS_PACIENTE = (
    "rut",
    "nombres",
    "apellidos",
    "fecha_nacimiento",
    "discapacidad",
    "embarazada",
    "postrado",
)
COLUMNAS_RECETA = (
    "folio",
    "fecha_emision",
    "vigencia_meses",
    "entregas_maximas",
    "entregas_realizadas",
)
COLUMNAS = COLUMNAS_PACIENTE + COLUMNAS_RECETA
COLUMNAS_OBLIGATORIAS = (
    "rut",
    "nombres",
    "apellidos",
    "fecha_nacimiento",
    "folio",
    "fecha_emision",
    "vigencia_meses",
    "entregas_maximas",
)
EXTENSIONES = (".csv", ".xlsx")
# Si alguna de estas columnas viene llena, la fila trae una receta y deben venir todas.
COLUMNAS_QUE_DEFINEN_RECETA = ("folio", "fecha_emision", "vigencia_meses", "entregas_maximas")
# Nombres de los campos de la fila ya normalizada (los de Paciente y RecetaCronica).
CAMPOS_PACIENTE = (
    "rut",
    "nombres",
    "apellidos",
    "fecha_nacimiento",
    "tiene_discapacidad",
    "esta_embarazada",
    "es_postrado",
)
CAMPOS_RECETA = ("folio", "fecha_emision", "vigencia_meses", "entregas_maximas", "entregas_realizadas")

LARGO_MAXIMO_NOMBRE = 100
LARGO_MAXIMO_FOLIO = 20
RANGO_VIGENCIA_MESES = (1, 24)
RANGO_ENTREGAS_MAXIMAS = (1, 24)
RANGO_ENTREGAS_REALIZADAS = (0, 24)
VERDADEROS = {"si", "sí", "1"}
FALSOS = {"no", "0", ""}
FORMATOS_FECHA = ("%d-%m-%Y", "%d/%m/%Y", "%Y-%m-%d")
TAMANO_LOTE_CONSULTA = 500  # RUT o folios por consulta `__in`, holgado bajo el límite de SQLite


class ErrorDeArchivo(Exception):
    """El archivo no se puede usar entero; el mensaje se muestra tal cual a la QF."""


class ErrorDeCarga(Exception):
    """Una acción que no corresponde sobre una carga; el mensaje se muestra tal cual a la QF."""


@dataclass(frozen=True)
class FilaLeida:
    numero: int  # como en Excel: el encabezado es la fila 1
    valores: dict  # columna conocida -> texto sin espacios a los lados ("" si viene vacía)


@dataclass(frozen=True)
class Lectura:
    encabezados: list
    filas: list


def normalizar_encabezado(texto):
    """`" Fecha_Emisión "` → `"fecha_emision"`: sin mayúsculas, tildes ni espacios a los lados."""
    descompuesto = unicodedata.normalize("NFKD", str(texto or "").strip().lower())
    return "".join(c for c in descompuesto if not unicodedata.combining(c))


def leer_archivo(archivo):
    """Lee un archivo subido (`.csv` o `.xlsx`) y devuelve sus encabezados y sus filas no vacías."""
    nombre = archivo.name or ""
    extension = Path(nombre).suffix.lower()
    if extension not in EXTENSIONES:
        raise ErrorDeArchivo("El archivo debe ser .csv o .xlsx.")

    maximo = settings.TURNOS_PADRON_BYTES_MAXIMOS
    contenido = archivo.read(maximo + 1)
    if len(contenido) > maximo:
        raise ErrorDeArchivo(f"El archivo pesa más de {maximo // (1024 * 1024)} MB.")
    if not contenido.strip():
        raise ErrorDeArchivo("El archivo está vacío.")

    filas_crudas = _filas_csv(contenido) if extension == ".csv" else _filas_xlsx(contenido)
    return _armar_lectura(filas_crudas)


def _filas_csv(contenido):
    try:
        texto = contenido.decode("utf-8-sig")
    except UnicodeDecodeError:
        try:
            texto = contenido.decode("cp1252")
        except UnicodeDecodeError:
            raise ErrorDeArchivo("No se pudo leer el archivo: la codificación no es UTF-8 ni Windows-1252.")

    primera_linea = next((linea for linea in texto.splitlines() if linea.strip()), "")
    separador = ";" if primera_linea.count(";") > primera_linea.count(",") else ","
    try:
        for numero, celdas in enumerate(csv.reader(io.StringIO(texto, newline=""), delimiter=separador), start=1):
            yield numero, celdas
    except csv.Error:
        raise ErrorDeArchivo("No se pudo leer el archivo: el CSV está mal formado.")


def _filas_xlsx(contenido):
    from openpyxl import load_workbook

    try:
        libro = load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
    except Exception:  # openpyxl lanza distintos errores ante un archivo dañado o que no es Excel
        raise ErrorDeArchivo("No se pudo leer el archivo: no es un Excel .xlsx válido.")
    try:
        if not libro.worksheets:
            raise ErrorDeArchivo("El archivo no tiene hojas.")
        hoja = libro.worksheets[0]
        for numero, celdas in enumerate(hoja.iter_rows(min_row=1, values_only=True), start=1):
            yield numero, [_texto_de_celda(celda) for celda in celdas]
    finally:
        libro.close()


def _texto_de_celda(valor):
    if valor is None:
        return ""
    if isinstance(valor, bool):
        return "1" if valor else "0"
    if isinstance(valor, datetime):
        return valor.date().isoformat()
    if isinstance(valor, date):
        return valor.isoformat()
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    return str(valor)


def _armar_lectura(filas_crudas):
    maximo = settings.TURNOS_PADRON_FILAS_MAXIMAS
    encabezados = None
    posiciones = {}
    filas = []
    for numero, celdas in filas_crudas:
        celdas = [str(celda).strip() for celda in celdas]
        if not any(celdas):
            continue
        if encabezados is None:
            encabezados = [normalizar_encabezado(celda) for celda in celdas]
            faltantes = [c for c in COLUMNAS_OBLIGATORIAS if c not in encabezados]
            if faltantes:
                raise ErrorDeArchivo("Faltan columnas obligatorias: " + ", ".join(faltantes) + ".")
            # Si una columna viene repetida, vale la primera.
            for posicion, encabezado in reversed(list(enumerate(encabezados))):
                if encabezado in COLUMNAS:
                    posiciones[encabezado] = posicion
            continue
        if len(filas) == maximo:
            raise ErrorDeArchivo(f"El archivo tiene más de {maximo} filas.")
        valores = {
            columna: celdas[posicion] if posicion < len(celdas) else ""
            for columna, posicion in posiciones.items()
        }
        for columna in COLUMNAS:
            valores.setdefault(columna, "")
        filas.append(FilaLeida(numero, valores))

    if encabezados is None:
        raise ErrorDeArchivo("El archivo está vacío.")
    return Lectura(encabezados, filas)


# --- Validación de filas ---------------------------------------------------------------------


@dataclass(frozen=True)
class Validacion:
    validas: list  # filas normalizadas, en el orden del archivo y listas para guardarse en JSON
    rechazos: list  # [{"fila": 12, "rut": "12.345.678-9", "motivo": "..."}], por número de fila


class _FilaInvalida(Exception):
    pass


def validar_filas(filas, hoy=None):
    """Valida las filas leídas sin consultar la base.

    Cada fila rechazada lleva un solo motivo, el primero que falla, en este orden: RUT, formato de
    los campos, receta incompleta, entregas realizadas sobre el máximo, un mismo RUT con datos de
    paciente distintos y un folio repetido en el archivo. Las dos últimas reglas comparan solo las
    filas que pasaron las cuatro primeras.
    """
    hoy = hoy or timezone.localdate()
    rechazos = {}
    normalizadas = []
    for fila in filas:
        try:
            normalizadas.append(_normalizar_fila(fila, hoy))
        except _FilaInvalida as error:
            rechazos[fila.numero] = _rechazo(fila.numero, fila.valores["rut"], str(error))

    # 5. Un mismo RUT con datos de paciente distintos: se rechazan todas sus filas.
    por_rut = defaultdict(list)
    for normalizada in normalizadas:
        por_rut[normalizada["rut"]].append(normalizada)
    for grupo in por_rut.values():
        if len({_datos_del_paciente(n) for n in grupo}) > 1:
            motivo = f"Datos del paciente distintos en las filas {_enumerar([n['fila'] for n in grupo])}."
            for normalizada in grupo:
                rechazos[normalizada["fila"]] = _rechazo(normalizada["fila"], normalizada["rut_original"], motivo)

    # 6. Un folio repetido en el archivo: se rechazan todas sus filas (salvo las que ya tienen motivo).
    por_folio = defaultdict(list)
    for normalizada in normalizadas:
        if normalizada["receta"]:
            por_folio[normalizada["receta"]["folio"]].append(normalizada)
    for folio, grupo in por_folio.items():
        if len(grupo) > 1:
            motivo = f"El folio {folio} se repite en las filas {_enumerar([n['fila'] for n in grupo])}."
            for normalizada in grupo:
                rechazos.setdefault(
                    normalizada["fila"], _rechazo(normalizada["fila"], normalizada["rut_original"], motivo)
                )

    validas = [n for n in normalizadas if n["fila"] not in rechazos]
    return Validacion(validas, [rechazos[numero] for numero in sorted(rechazos)])


def _rechazo(numero, rut, motivo):
    return {"fila": numero, "rut": rut, "motivo": motivo}


def _datos_del_paciente(normalizada):
    return tuple(normalizada[campo] for campo in CAMPOS_PACIENTE)


def _enumerar(elementos):
    """`[12, 40, 55]` → `"12, 40 y 55"`."""
    textos = [str(elemento) for elemento in elementos]
    return textos[0] if len(textos) == 1 else ", ".join(textos[:-1]) + " y " + textos[-1]


def _normalizar_fila(fila, hoy):
    """Revisa las reglas 1 a 4 de una fila y la devuelve normalizada, o lanza `_FilaInvalida`."""
    valores = fila.valores

    # 1. RUT
    if not valores["rut"]:
        raise _FilaInvalida("Falta el RUT.")
    try:
        rut = normalizar_rut(valores["rut"])
        validar_rut(rut)
    except ValidationError as error:
        raise _FilaInvalida(error.messages[0])

    # 2. Formato de los campos (los de la receta solo si vienen llenos: los vacíos son la regla 3)
    normalizada = {
        "fila": fila.numero,
        "rut_original": valores["rut"],
        "rut": rut,
        "nombres": _texto(valores, "nombres", LARGO_MAXIMO_NOMBRE),
        "apellidos": _texto(valores, "apellidos", LARGO_MAXIMO_NOMBRE),
        "fecha_nacimiento": _fecha(valores, "fecha_nacimiento", hoy, obligatoria=True),
        "tiene_discapacidad": _bandera(valores, "discapacidad"),
        "esta_embarazada": _bandera(valores, "embarazada"),
        "es_postrado": _bandera(valores, "postrado"),
    }
    folio = valores["folio"]
    if folio and (len(folio) > LARGO_MAXIMO_FOLIO or re.search(r"\s", folio)):
        raise _FilaInvalida(f"El folio debe tener hasta {LARGO_MAXIMO_FOLIO} caracteres, sin espacios.")
    fecha_emision = _fecha(valores, "fecha_emision", hoy)
    vigencia = _entero(valores, "vigencia_meses", *RANGO_VIGENCIA_MESES)
    maximas = _entero(valores, "entregas_maximas", *RANGO_ENTREGAS_MAXIMAS)
    realizadas = _entero(valores, "entregas_realizadas", *RANGO_ENTREGAS_REALIZADAS)

    # 3. Receta incompleta
    trae_receta = any(valores[columna] for columna in COLUMNAS_RECETA)
    if trae_receta:
        faltantes = [columna for columna in COLUMNAS_QUE_DEFINEN_RECETA if not valores[columna]]
        if faltantes:
            raise _FilaInvalida(f"Receta incompleta: falta {_enumerar(faltantes)}.")

    # 4. Entregas realizadas sobre el máximo
    if trae_receta and realizadas is not None and realizadas > maximas:
        raise _FilaInvalida(
            f"Las entregas realizadas ({realizadas}) superan las entregas máximas ({maximas})."
        )

    normalizada["receta"] = (
        {
            "folio": folio,
            "fecha_emision": fecha_emision,
            "vigencia_meses": vigencia,
            "entregas_maximas": maximas,
            "entregas_realizadas": realizadas or 0,
        }
        if trae_receta
        else None
    )
    return normalizada


def _texto(valores, columna, largo_maximo):
    texto = valores[columna]
    if not texto:
        raise _FilaInvalida(f"Falta el campo {columna}.")
    if len(texto) > largo_maximo:
        raise _FilaInvalida(f"El campo {columna} supera los {largo_maximo} caracteres.")
    return texto


def _fecha(valores, columna, hoy, obligatoria=False):
    """Devuelve la fecha en ISO (`aaaa-mm-dd`), o None si viene vacía y no es obligatoria."""
    texto = valores[columna]
    if not texto:
        if obligatoria:
            raise _FilaInvalida(f"Falta el campo {columna}.")
        return None
    for formato in FORMATOS_FECHA:
        try:
            fecha = datetime.strptime(texto, formato).date()
            break
        except ValueError:
            continue
    else:
        raise _FilaInvalida(f"El campo {columna} no es una fecha válida: use dd-mm-aaaa.")
    if fecha > hoy:
        raise _FilaInvalida(f"El campo {columna} no puede ser una fecha futura.")
    return fecha.isoformat()


def _bandera(valores, columna):
    texto = valores[columna].lower()
    if texto in VERDADEROS:
        return True
    if texto in FALSOS:
        return False
    raise _FilaInvalida(f"El campo {columna} debe ser sí o no.")


def _entero(valores, columna, minimo, maximo):
    """Devuelve el entero, o None si viene vacío (si falta en una receta, lo dice la regla 3)."""
    texto = valores[columna]
    if not texto:
        return None
    if not re.fullmatch(r"[0-9]+", texto) or not minimo <= int(texto) <= maximo:
        raise _FilaInvalida(f"El campo {columna} debe ser un número entero de {minimo} a {maximo}.")
    return int(texto)


# --- Comparación con la base y vista previa --------------------------------------------------


@dataclass
class Plan:
    """Lo que haría la carga sobre la base actual. Lo arma `planificar`, que no escribe nada."""

    filas: list  # filas válidas que siguen tras comparar con la base
    rechazos: list  # solo los de la regla 7 (folio de otro RUT)
    advertencias: list
    pacientes_existentes: dict  # rut -> Paciente
    recetas_existentes: dict  # folio -> RecetaCronica
    pacientes_creados: set
    pacientes_actualizados: set
    pacientes_sin_cambios: set
    recetas_creadas: set
    recetas_actualizadas: set
    recetas_sin_cambios: set

    def conteos(self):
        return {
            "pacientes_creados": len(self.pacientes_creados),
            "pacientes_actualizados": len(self.pacientes_actualizados),
            "pacientes_sin_cambios": len(self.pacientes_sin_cambios),
            "recetas_creadas": len(self.recetas_creadas),
            "recetas_actualizadas": len(self.recetas_actualizadas),
            "recetas_sin_cambios": len(self.recetas_sin_cambios),
        }


def planificar(filas):
    """Compara las filas válidas con la base, con consultas por lotes, sin escribir nada.

    Rechaza la regla 7 (el folio ya existe para otro RUT), clasifica cada paciente (por RUT) y
    cada receta (por folio) en creado, actualizado o sin cambios, y anota una advertencia cuando
    la base registra más entregas que el archivo: ese contador nunca baja.
    """
    pacientes = {
        p.rut: p for p in _por_lotes(Paciente.objects.all(), "rut__in", {f["rut"] for f in filas})
    }
    folios = {f["receta"]["folio"] for f in filas if f["receta"]}
    recetas = {
        r.folio: r
        for r in _por_lotes(RecetaCronica.objects.select_related("paciente"), "folio__in", folios)
    }

    plan = Plan([], [], [], pacientes, recetas, set(), set(), set(), set(), set(), set())
    for fila in filas:
        receta = fila["receta"]
        existente = recetas.get(receta["folio"]) if receta else None
        if existente and existente.paciente.rut != fila["rut"]:
            plan.rechazos.append(_rechazo(
                fila["fila"], fila["rut_original"], f"El folio {receta['folio']} ya está registrado para otro RUT."
            ))
            continue
        plan.filas.append(fila)

    for fila in plan.filas:
        rut = fila["rut"]
        if rut in plan.pacientes_creados | plan.pacientes_actualizados | plan.pacientes_sin_cambios:
            continue  # otra fila del mismo paciente; validar_filas ya comprobó que trae los mismos datos
        paciente = pacientes.get(rut)
        if paciente is None:
            plan.pacientes_creados.add(rut)
        elif _cambios_del_paciente(paciente, fila):
            plan.pacientes_actualizados.add(rut)
        else:
            plan.pacientes_sin_cambios.add(rut)

    for fila in plan.filas:
        receta = fila["receta"]
        if not receta:
            continue
        existente = recetas.get(receta["folio"])
        if existente is None:
            plan.recetas_creadas.add(receta["folio"])
            continue
        if existente.entregas_realizadas > receta["entregas_realizadas"]:
            plan.advertencias.append({
                "fila": fila["fila"],
                "folio": receta["folio"],
                "mensaje": (
                    f"El sistema registra {existente.entregas_realizadas} entregas y el archivo trae "
                    f"{receta['entregas_realizadas']}: se conservan {existente.entregas_realizadas}."
                ),
            })
        if _cambios_de_la_receta(existente, receta):
            plan.recetas_actualizadas.add(receta["folio"])
        else:
            plan.recetas_sin_cambios.add(receta["folio"])
    return plan


def _por_lotes(consulta, filtro, valores):
    valores = sorted(valores)
    for inicio in range(0, len(valores), TAMANO_LOTE_CONSULTA):
        yield from consulta.filter(**{filtro: valores[inicio:inicio + TAMANO_LOTE_CONSULTA]})


def valores_del_paciente(fila):
    """Los campos de `Paciente` que trae una fila normalizada, con la fecha como `date`."""
    valores = {campo: fila[campo] for campo in CAMPOS_PACIENTE if campo != "rut"}
    valores["fecha_nacimiento"] = date.fromisoformat(valores["fecha_nacimiento"])
    return valores


def valores_de_la_receta(receta, existente=None):
    """Los campos de `RecetaCronica` que quedarían; el contador de entregas nunca baja."""
    valores = {campo: receta[campo] for campo in CAMPOS_RECETA if campo != "folio"}
    valores["fecha_emision"] = date.fromisoformat(valores["fecha_emision"])
    if existente is not None:
        valores["entregas_realizadas"] = max(existente.entregas_realizadas, receta["entregas_realizadas"])
    return valores


def _cambios_del_paciente(paciente, fila):
    return {
        campo: valor
        for campo, valor in valores_del_paciente(fila).items()
        if getattr(paciente, campo) != valor
    }


def _cambios_de_la_receta(existente, receta):
    return {
        campo: valor
        for campo, valor in valores_de_la_receta(receta, existente).items()
        if getattr(existente, campo) != valor
    }


def preparar_carga(archivo, usuario):
    """Lee y valida el archivo, lo compara con la base y crea la carga pendiente (vista previa).

    No escribe pacientes ni recetas, ni guarda el archivo: solo su nombre, tamaño y huella. Lanza
    `ErrorDeArchivo` si el archivo no se puede usar entero, y en ese caso no crea la carga.
    """
    archivo.seek(0)
    lectura = leer_archivo(archivo)
    archivo.seek(0)
    huella = hashlib.sha256()
    for trozo in archivo.chunks():
        huella.update(trozo)

    validacion = validar_filas(lectura.filas)
    plan = planificar(validacion.validas)
    rechazos = sorted(validacion.rechazos + plan.rechazos, key=lambda rechazo: rechazo["fila"])
    return CargaPadron.objects.create(
        usuario=usuario,
        estado=EstadoCarga.PENDIENTE,
        nombre_archivo=Path(archivo.name).name[:255],
        bytes_archivo=archivo.size,
        huella_archivo=huella.hexdigest(),
        filas_totales=len(lectura.filas),
        filas_rechazadas=len(rechazos),
        filas_validas=plan.filas,
        rechazos=rechazos,
        advertencias=plan.advertencias,
        **plan.conteos(),
    )


# --- Confirmación, cancelación y descarte ----------------------------------------------------


def _limite_pendiente():
    return timezone.now() - timedelta(minutes=settings.TURNOS_PADRON_MINUTOS_PENDIENTE)


def descartar_cargas_vencidas():
    """Descarta las cargas pendientes que pasaron el plazo sin confirmarse. Devuelve cuántas."""
    return CargaPadron.objects.filter(
        estado=EstadoCarga.PENDIENTE, creada_en__lt=_limite_pendiente()
    ).update(estado=EstadoCarga.DESCARTADA, filas_validas=None)


def cancelar_carga(carga):
    """Cancela una carga pendiente y borra sus filas; las demás no se pueden cancelar."""
    cancelada = CargaPadron.objects.filter(pk=carga.pk, estado=EstadoCarga.PENDIENTE).update(
        estado=EstadoCarga.CANCELADA, filas_validas=None
    )
    if not cancelada:
        raise ErrorDeCarga("Esta carga ya no está pendiente: no se puede cancelar.")
    carga.refresh_from_db()
    return carga


def confirmar_carga(carga, usuario):
    """Aplica una carga pendiente: crea y actualiza pacientes y recetas en una sola transacción.

    La carga se toma con un UPDATE condicional sobre `estado = 'pendiente'`, así dos
    confirmaciones a la vez la aplican una sola vez. Antes de escribir se vuelve a comparar con la
    base, porque pudo cambiar desde la vista previa. Si algo falla, no queda nada escrito y la
    carga vuelve a `pendiente`. `usuario` es quien confirma; la bitácora guarda quién la subió.
    """
    if CargaPadron.objects.filter(
        pk=carga.pk, estado=EstadoCarga.PENDIENTE, creada_en__lt=_limite_pendiente()
    ).update(estado=EstadoCarga.DESCARTADA, filas_validas=None):
        raise ErrorDeCarga(
            f"La carga pasó {settings.TURNOS_PADRON_MINUTOS_PENDIENTE} minutos sin confirmarse y "
            "se descartó. Suba el archivo de nuevo."
        )
    tomada = CargaPadron.objects.filter(pk=carga.pk, estado=EstadoCarga.PENDIENTE).update(
        estado=EstadoCarga.APLICANDO
    )
    if not tomada:
        raise ErrorDeCarga("Esta carga ya no está pendiente: no se puede aplicar.")

    try:
        with transaction.atomic():
            carga.refresh_from_db()
            _aplicar(carga)
    except Exception:
        CargaPadron.objects.filter(pk=carga.pk, estado=EstadoCarga.APLICANDO).update(
            estado=EstadoCarga.PENDIENTE
        )
        raise
    return carga


def _aplicar(carga):
    plan = planificar(carga.filas_validas)
    primera_fila = {}
    for fila in plan.filas:
        primera_fila.setdefault(fila["rut"], fila)

    Paciente.objects.bulk_create(
        [
            Paciente(rut=rut, carga=carga, **valores_del_paciente(primera_fila[rut]))
            for rut in sorted(plan.pacientes_creados)
        ],
        batch_size=TAMANO_LOTE_CONSULTA,
    )
    actualizados = []
    for rut in sorted(plan.pacientes_actualizados):
        paciente = plan.pacientes_existentes[rut]
        for campo, valor in valores_del_paciente(primera_fila[rut]).items():
            setattr(paciente, campo, valor)
        paciente.carga = carga
        actualizados.append(paciente)
    Paciente.objects.bulk_update(
        actualizados,
        [campo for campo in CAMPOS_PACIENTE if campo != "rut"] + ["carga"],
        batch_size=TAMANO_LOTE_CONSULTA,
    )

    # Los pacientes recién creados se vuelven a leer para tener su clave primaria.
    pacientes = {
        p.rut: p
        for p in _por_lotes(Paciente.objects.all(), "rut__in", {f["rut"] for f in plan.filas if f["receta"]})
    }
    nuevas, cambiadas = [], []
    for fila in plan.filas:
        receta = fila["receta"]
        if not receta:
            continue
        if receta["folio"] in plan.recetas_creadas:
            nuevas.append(RecetaCronica(
                paciente=pacientes[fila["rut"]],
                folio=receta["folio"],
                carga=carga,
                **valores_de_la_receta(receta),
            ))
        elif receta["folio"] in plan.recetas_actualizadas:
            existente = plan.recetas_existentes[receta["folio"]]
            for campo, valor in valores_de_la_receta(receta, existente).items():
                setattr(existente, campo, valor)
            # El contador se calcula en la base: si una atención sumó una entrega desde que se
            # leyó la receta, el máximo la respeta y el contador nunca baja.
            existente.entregas_realizadas = Greatest(
                F("entregas_realizadas"), Value(receta["entregas_realizadas"])
            )
            existente.carga = carga
            cambiadas.append(existente)
    RecetaCronica.objects.bulk_create(nuevas, batch_size=TAMANO_LOTE_CONSULTA)
    RecetaCronica.objects.bulk_update(
        cambiadas,
        [campo for campo in CAMPOS_RECETA if campo != "folio"] + ["carga"],
        batch_size=TAMANO_LOTE_CONSULTA,
    )

    rechazos = sorted(carga.rechazos + plan.rechazos, key=lambda rechazo: rechazo["fila"])
    for campo, valor in plan.conteos().items():
        setattr(carga, campo, valor)
    carga.rechazos = rechazos
    carga.filas_rechazadas = len(rechazos)
    carga.advertencias = plan.advertencias
    carga.filas_validas = None
    carga.estado = EstadoCarga.APLICADA
    carga.aplicada_en = timezone.now()
    carga.save()


# --- Archivos que se descargan ---------------------------------------------------------------

# Una fila sintética: subida sin cambios, la plantilla produce una carga sin rechazos.
FILA_EJEMPLO_PLANTILLA = {
    "rut": "11.111.111-1",
    "nombres": "Ejemplo",
    "apellidos": "Sintético Prueba",
    "fecha_nacimiento": "15-03-1944",
    "discapacidad": "no",
    "embarazada": "no",
    "postrado": "si",
    "folio": "EJEMPLO-0001",
    "fecha_emision": "01-01-2026",
    "vigencia_meses": "12",
    "entregas_maximas": "6",
    "entregas_realizadas": "0",
}
INICIOS_DE_FORMULA = ("=", "+", "-", "@")


def neutralizar_celda(valor):
    """Antepone `'` a una celda que Excel tomaría como fórmula (inyección de fórmulas en CSV)."""
    texto = str(valor)
    return "'" + texto if texto.startswith(INICIOS_DE_FORMULA) else texto


def csv_para_excel(filas):
    """CSV con `;` y en UTF-8 con BOM, para que Excel en español lo abra con columnas y tildes."""
    salida = io.StringIO()
    escritor = csv.writer(salida, delimiter=";", lineterminator="\r\n")
    for fila in filas:
        escritor.writerow([neutralizar_celda(valor) for valor in fila])
    return "﻿" + salida.getvalue()


def csv_de_plantilla():
    return csv_para_excel([COLUMNAS, [FILA_EJEMPLO_PLANTILLA[columna] for columna in COLUMNAS]])


def csv_de_rechazos(carga):
    filas = [(r["fila"], r["rut"], r["motivo"]) for r in carga.rechazos]
    return csv_para_excel([("fila", "rut", "motivo"), *filas])
